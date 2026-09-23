"""Keepa watch lists: the exact 'needed move' maths, the half-threshold guarantee (a product that becomes a lead through
ANY mix of UK rise and EU fall trips at least one alert), bucketing, the selection rules and the files written.
Isolated in-memory SQLite and a temp folder -- no live database access, no Keepa or SP-API calls."""
import os
import tempfile
import unittest
from datetime import datetime, timedelta, timezone
from types import SimpleNamespace
from unittest.mock import patch

from openpyxl import load_workbook
from sqlalchemy import create_engine
from sqlalchemy.orm import sessionmaker
from sqlalchemy.pool import StaticPool

from app.database.base import Base
from app.database.models import ListingRestriction, ProductRecord, WatchedProduct
from app.services import keepa_watch_list_service as kw
from app.services.fee_engine import FeeEngine
from app.services.keepa_watch_list_service import (
    ALERT_BUCKETS, MAX_NEEDED_MOVE_PCT, TARGET_ROI_PCT, KeepaWatchListService, alert_bucket, eu_cost_for_roi,
    needed_moves, roi_fn_for, uk_price_for_roi, _stored_rate_roi,
)

NOW = datetime.now(timezone.utc).replace(tzinfo=None)


def flat_model(uk_vat=0.2, rate=0.15, fba=3.5, eu_vat=0.20):
    """A flat-referral fee model: profit is exactly linear in both prices, so the half-threshold guarantee is exact."""
    return _stored_rate_roi(uk_vat, rate, fba, eu_vat)


def category_model(category="toys & games", fba=3.5, eu_vat=0.20):
    return lambda price, cost: FeeEngine.roi_at_price(price, cost, category, fba, eu_vat)


class NeededMoveMathsTests(unittest.TestCase):
    def test_the_cost_that_hits_the_target_really_does(self):
        for roi_fn in (flat_model(), category_model()):
            for price, cost in ((40, 30), (25, 18), (120, 90), (12, 8)):
                target_cost = eu_cost_for_roi(roi_fn, price, cost_now=cost)
                self.assertAlmostEqual(roi_fn(price, target_cost), TARGET_ROI_PCT, delta=0.15, msg=(price, cost))

    def test_the_price_that_hits_the_target_really_does(self):
        for roi_fn in (flat_model(), category_model()):
            for price, cost in ((32, 24), (25, 20), (80, 64)):
                target_price = uk_price_for_roi(roi_fn, cost, price)
                self.assertAlmostEqual(roi_fn(target_price, cost), TARGET_ROI_PCT, delta=0.3, msg=(price, cost))
                self.assertGreater(target_price, price)

    def test_direction_of_both_moves(self):
        roi_fn = flat_model()
        moves = needed_moves(roi_fn, 32.0, 24.0)
        self.assertLess(moves["roi_now"], TARGET_ROI_PCT)
        self.assertLess(moves["eu_target_cost"], 24.0)            # the EU cost has to come DOWN
        self.assertGreater(moves["uk_target_price"], 32.0)        # the UK price has to go UP
        self.assertGreater(moves["eu_drop_pct"], 0)
        self.assertGreater(moves["uk_rise_pct"], 0)

    def test_already_a_lead_needs_nothing(self):
        moves = needed_moves(flat_model(), 40.0, 15.0)
        self.assertGreaterEqual(moves["roi_now"], TARGET_ROI_PCT)
        self.assertEqual((moves["eu_drop_pct"], moves["uk_rise_pct"]), (0.0, 0.0))

    def test_unreachable_sides_are_none_not_garbage(self):
        # Fees swallow a GBP 2 price: no EU cost, however low, gets 25%.
        self.assertIsNone(eu_cost_for_roi(flat_model(), 2.0, cost_now=1.5))
        # A cost of 100 against a price of 10 can't be fixed by any rise within 3x.
        self.assertIsNone(uk_price_for_roi(flat_model(), 100.0, 10.0))
        moves = needed_moves(flat_model(), 10.0, 100.0)
        self.assertIsNone(moves["uk_rise_pct"])

    def test_the_uk_also_fell_example_from_the_conversation(self):
        """DE down 10% and UK down 20% (22.1% -> 6.3% ROI, with the Digital Services Fee and UK-VAT netting). What would fix it? The UK recovering."""
        roi_fn = category_model()
        before = roi_fn(40.0, 22.0)
        after = needed_moves(roi_fn, 32.0, 19.8)
        self.assertAlmostEqual(before, 22.1, delta=0.2)
        self.assertAlmostEqual(after["roi_now"], 6.3, delta=0.3)
        # Back at 40 the same EU cost is 33%, so the UK only has to recover part of the way.
        self.assertLess(after["uk_rise_pct"], 25.0)
        self.assertGreater(roi_fn(40.0, 19.8), TARGET_ROI_PCT)


class HalfThresholdGuaranteeTests(unittest.TestCase):
    """The reason for two lists: any mix (UK rise u, EU fall d) that makes a lead trips at least one half-threshold."""

    def combos(self, roi_fn, price, cost):
        moves = needed_moves(roi_fn, price, cost)
        s_eu, s_uk = moves["eu_drop_pct"] / 100, moves["uk_rise_pct"] / 100
        return s_eu, s_uk

    def test_the_lead_boundary_is_a_straight_line_in_the_two_moves(self):
        roi_fn = flat_model()
        price, cost = 32.0, 24.0
        s_eu, s_uk = self.combos(roi_fn, price, cost)
        for share in (0.0, 0.25, 0.5, 0.8, 1.0):                   # u/s_uk + d/s_eu == 1 all along the boundary
            u, d = share * s_uk, (1 - share) * s_eu
            self.assertAlmostEqual(roi_fn(price * (1 + u), cost * (1 - d)), TARGET_ROI_PCT, delta=0.3, msg=share)

    def test_every_lead_making_mix_trips_at_least_one_half_threshold(self):
        checked = 0
        for roi_fn in (flat_model(), flat_model(rate=0.08, fba=2.0, eu_vat=0.22)):
            for price, cost in ((32, 24), (45, 34), (25, 19.5), (90, 70)):
                s_eu, s_uk = self.combos(roi_fn, price, cost)
                for i in range(0, 41):
                    for j in range(0, 41):
                        u, d = s_uk * i / 20, s_eu * j / 20                          # mixes up to 2x each side
                        if roi_fn(price * (1 + u), cost * (1 - d)) >= TARGET_ROI_PCT + 0.05:   # a genuine lead
                            checked += 1
                            self.assertTrue(u >= s_uk / 2 - 1e-9 or d >= s_eu / 2 - 1e-9, (price, cost, u, d))
        self.assertGreater(checked, 500)                                             # the property was really exercised

    def test_a_single_side_alert_alone_would_have_missed_some(self):
        """The point of the second list: watching only the EU price at its full threshold misses combined moves."""
        roi_fn = flat_model()
        price, cost = 32.0, 24.0
        s_eu, s_uk = self.combos(roi_fn, price, cost)
        u, d = 0.6 * s_uk, 0.6 * s_eu                                               # both partly moved: a lead
        self.assertGreaterEqual(roi_fn(price * (1 + u), cost * (1 - d)), TARGET_ROI_PCT)
        self.assertLess(d, s_eu)                                                    # ...but the EU-only alert (full threshold) never fired
        self.assertGreaterEqual(d, s_eu / 2)                                        # while the half-threshold alert does


class AlertBucketTests(unittest.TestCase):
    def test_half_of_the_needed_move_rounded_down_to_a_bucket(self):
        self.assertEqual(alert_bucket(40), 20)      # half 20
        self.assertEqual(alert_bucket(30), 15)
        self.assertEqual(alert_bucket(29.9), 10)    # half 14.95 -> rounds DOWN, never up
        self.assertEqual(alert_bucket(20), 10)
        self.assertEqual(alert_bucket(12), 5)       # half 6
        self.assertEqual(alert_bucket(10), 5)
        self.assertEqual(alert_bucket(6), 3)        # half 3
        self.assertEqual(alert_bucket(3), 3)        # half 1.5: floored at the smallest bucket
        self.assertEqual(alert_bucket(0.5), 3)

    def test_a_bucket_never_exceeds_half_the_need_except_at_the_floor(self):
        for needed in [x / 10 for x in range(5, int(MAX_NEEDED_MOVE_PCT * 10) + 1)]:
            bucket = alert_bucket(needed)
            self.assertIn(bucket, ALERT_BUCKETS)
            self.assertTrue(bucket <= needed / 2 or bucket == ALERT_BUCKETS[0], needed)

    def test_no_alert_when_there_is_nothing_sensible_to_set(self):
        self.assertIsNone(alert_bucket(None))
        self.assertIsNone(alert_bucket(0))
        self.assertIsNone(alert_bucket(-3))
        self.assertIsNone(alert_bucket(MAX_NEEDED_MOVE_PCT + 0.1))


def record(asin="B0AAAAAAAA", price=32.0, cost=24.0, rate=0.15, **kw):
    """A ProductRecord whose stored ROI was computed with the same fee model the service will pick."""
    model = flat_model(rate=rate)
    fields = dict(asin=asin, title=f"Product {asin}", brand="Brand", category_name="", buy_box_now=price,
                  buy_box_90d=price * 1.1, best_source_marketplace="DE", best_source_cost_gbp=cost, fba_fee=3.5,
                  referral_rate_used=rate, uk_vat_rate_used=0.2, eu_vat_rate_used=0.20, roi=round(model(price, cost), 2) if cost > 0 else 0.0,
                  roi_90d=0.0, monthly_sales=50, sales_drops_30d=0, recommendation="IGNORE", scanned_at=NOW - timedelta(days=2))
    fields.update(kw)
    return ProductRecord(**fields)


class DigitalServicesFeeInTheWatchListModelTests(unittest.TestCase):
    """The watch list keeps its own copy of the profit formula (the stored-rate model) -- it must charge the same
    Digital Services Fee as FeeEngine, or its "needed move" maths would disagree with what Atlas shows."""

    def test_the_stored_rate_model_charges_the_fee_like_fee_engine_does(self):
        # Toys & Games is a flat 15% referral, so the two models must agree.
        stored = _stored_rate_roi(0.2, 0.15, 3.59, 0.22)
        for price, cost in ((43.51, 25.68), (30.0, 12.0), (80.0, 40.0)):
            self.assertAlmostEqual(stored(price, cost), FeeEngine.roi_at_price(price, cost, "toys & games", 3.59, 0.22), delta=0.05)

    def test_switching_the_fee_off_gives_the_old_numbers(self):
        with_fee, without = _stored_rate_roi(0.2, 0.15, 3.59, 0.22), _stored_rate_roi(0.2, 0.15, 3.59, 0.22, include_dsf=False)
        referral = 43.51 * 0.15
        dsf = round(0.02 * (referral + 3.59), 2)
        self.assertAlmostEqual(without(43.51, 25.68) - with_fee(43.51, 25.68), dsf / 25.68 * 100, delta=0.02)


class RecordsScannedBeforeTheFeeStillValidateTests(unittest.TestCase):
    def test_a_cheap_record_priced_without_the_fee_is_still_trusted_and_the_model_returned_includes_it(self):
        # cost 3: the fee is ~3.7 ROI points, more than the 3-point mismatch bar -- checking only WITH the fee would
        # wrongly call this record's fee data unusable.
        old_way, new_way = _stored_rate_roi(0.2, 0.15, 3.5, 0.20, include_dsf=False), _stored_rate_roi(0.2, 0.15, 3.5, 0.20)
        r = record("B0CHEAPOLD", price=12.0, cost=3.0, roi=round(old_way(12.0, 3.0), 2))
        self.assertGreater(abs(new_way(12.0, 3.0) - r.roi), kw.ROI_MISMATCH_POINTS)       # the trap this guards against
        fn, gap = roi_fn_for(r)
        self.assertLess(gap, kw.ROI_MISMATCH_POINTS)
        self.assertLess(fn(12.0, 3.0), r.roi)                               # ...but what it returns DOES charge the fee

    def test_a_record_priced_with_the_fee_is_trusted_too(self):
        new_way = _stored_rate_roi(0.2, 0.15, 3.5, 0.20)
        r = record("B0CHEAPNEW", price=12.0, cost=3.0, roi=round(new_way(12.0, 3.0), 2))
        self.assertLess(roi_fn_for(r)[1], 0.5)


class RecordsPricedUnderTheOldLocalVatRulesTests(unittest.TestCase):
    """Until 2026-09-21 an Italian cost was netted by 22% VAT, a German one by 19%. Tamara reclaims UK VAT (20%) only, so
    every EU cost is now netted by 20% -- but a record scanned before that still stores its old ROI and old rate."""

    def old_rules_record(self, local_vat, price=32.0, cost=17.4):
        old = _stored_rate_roi(0.2, 0.15, 3.5, local_vat, include_dsf=False)
        return record("B0OLDLOCAL", price=price, cost=cost, eu_vat_rate_used=local_vat, roi=round(old(price, cost), 2))

    def test_an_italian_record_priced_at_22_percent_is_still_trusted(self):
        for local_vat in (0.19, 0.21, 0.22):
            r = self.old_rules_record(local_vat)
            fn, gap = roi_fn_for(r)
            self.assertLess(gap, 0.5, local_vat)

    def test_the_model_returned_nets_the_cost_by_uk_vat_whatever_the_record_stored(self):
        r = self.old_rules_record(0.22)
        fn, _ = roi_fn_for(r)
        now_rules = _stored_rate_roi(0.2, 0.15, 3.5, 0.20)
        self.assertAlmostEqual(fn(32.0, 17.4), now_rules(32.0, 17.4), places=6)
        self.assertLess(fn(32.0, 17.4), r.roi)                      # 20% nets less VAT off than 22%, and the fee is charged

    def test_a_german_record_priced_at_19_percent_comes_out_higher_on_the_new_rules_vat_alone(self):
        vat_only_new = _stored_rate_roi(0.2, 0.15, 3.5, 0.20, include_dsf=False)
        r = self.old_rules_record(0.19)
        self.assertGreater(vat_only_new(32.0, 17.4), r.roi)          # 20% removes MORE VAT than 19%: net cost is lower


class RoiModelChoiceTests(unittest.TestCase):
    def test_picks_the_model_that_reproduces_the_stored_roi(self):
        stored_rate = record(category_name="", rate=0.08)                         # blank category, real rate stored
        fn, gap = roi_fn_for(stored_rate)
        self.assertLess(gap, 0.5)
        self.assertAlmostEqual(fn(32.0, 24.0), stored_rate.roi, delta=0.5)

    def test_prefers_the_category_model_when_it_is_the_one_that_matches(self):
        r = record(category_name="Toys & Games", referral_rate_used=0.99)         # a nonsense stored rate
        r.roi = round(FeeEngine.roi_at_price(32.0, 24.0, "Toys & Games", 3.5, 0.20), 2)
        _, gap = roi_fn_for(r)
        self.assertLess(gap, 0.5)

    def test_the_stored_rate_model_wins_when_it_is_the_better_fit_even_with_a_category_present(self):
        r = record(category_name="Toys & Games", rate=0.08)          # ROI stored with an 8% referral rate, not the category's
        category_gap = abs(FeeEngine.roi_at_price(32.0, 24.0, "Toys & Games", 3.5, 0.20) - r.roi)
        self.assertGreater(category_gap, 3.0)                       # the category model alone would be well off
        fn, gap = roi_fn_for(r)
        self.assertLess(gap, 0.5)
        self.assertAlmostEqual(fn(32.0, 24.0), r.roi, delta=0.5)

    def test_a_record_with_nothing_to_go_on_reports_a_big_gap(self):
        r = record(category_name="", referral_rate_used=0.0)
        r.roi = 21.0                                                              # stored ROI the default model can't reproduce
        _, gap = roi_fn_for(r)
        self.assertGreater(gap, kw.ROI_MISMATCH_POINTS)


class _Db(unittest.TestCase):
    def setUp(self):
        self.engine = create_engine("sqlite://", poolclass=StaticPool, connect_args={"check_same_thread": False})
        Base.metadata.create_all(self.engine)
        self.sessions = sessionmaker(bind=self.engine)
        patcher = patch.object(kw, "SessionLocal", self.sessions)
        patcher.start()
        self.addCleanup(patcher.stop)
        self.addCleanup(self.engine.dispose)

    def add(self, rec, watched=False, restricted=False):
        with self.sessions() as db:
            db.add(rec)
            if watched:
                db.add(WatchedProduct(asin=rec.asin, title=rec.title))
            if restricted:
                db.add(ListingRestriction(asin=rec.asin, marketplace="UK", restricted=True))
            db.commit()

    def build(self):
        result = KeepaWatchListService.build(now=NOW)
        return {r["asin"]: r for r in result["rows"]}, {e["asin"]: e["reason"] for e in result["excluded"]}, result


class SelectionTests(_Db):
    def test_a_near_miss_gets_both_sides_and_a_sensible_bucket(self):
        self.add(record("B0NEARMISS", price=32.0, cost=17.4, roi_90d=0))               # ~19% ROI: a near miss
        rows, excluded, _ = self.build()
        row = rows["B0NEARMISS"]
        self.assertEqual(row["scopes"], "Near miss")
        self.assertLess(row["roi_now"], TARGET_ROI_PCT)
        self.assertGreater(row["eu_drop_needed_pct"], 0)
        self.assertGreater(row["uk_rise_needed_pct"], 0)
        self.assertIn(row["eu_alert_pct"], ALERT_BUCKETS)
        self.assertIn(row["uk_alert_pct"], ALERT_BUCKETS)
        self.assertLess(row["eu_target_cost_gbp"], row["eu_cost_gbp"])
        self.assertGreater(row["uk_target_price_gbp"], row["uk_now"])

    def test_the_uk_recovery_case_is_included_and_labelled(self):
        # UK 32 vs typical 40, EU 19.8: ~6% today, ~34% at 40.
        self.add(record("B0RECOVERY", price=32.0, cost=19.8, buy_box_90d=40.0, roi_90d=34.0, roi=6.5, monthly_sales=80))
        rows, excluded, _ = self.build()
        self.assertNotIn("B0RECOVERY", excluded)
        self.assertIn("UK recovery", rows["B0RECOVERY"]["scopes"])

    def test_watchlist_items_are_included_wherever_they_came_from(self):
        self.add(record("B0WATCHED0", price=32.0, cost=26.0, monthly_sales=20), watched=True)
        rows, _, _ = self.build()
        self.assertEqual(rows["B0WATCHED0"]["scopes"], "Watchlist")

    def test_each_exclusion_says_why(self):
        self.add(record("B0NOSOURCE", cost=0.0, roi=15.0), watched=True)
        self.add(record("B0NOSALES0", monthly_sales=0, sales_drops_30d=0, roi=15.0), watched=True)
        self.add(record("B0GATEDREC", roi=15.0, recommendation="GATED"), watched=True)
        self.add(record("B0RESTRICT", roi=15.0), watched=True, restricted=True)
        self.add(record("B0PLUGITEM", roi=15.0, title="Shark Cordless Vacuum Cleaner"), watched=True)
        self.add(record("B0STALEREC", roi=15.0, scanned_at=NOW - timedelta(days=200)), watched=True)
        self.add(record("B0ALREADYA", price=40.0, cost=15.0), watched=True)                        # ROI 100%+: a lead today
        self.add(record("B0TOOFAR00", price=12.0, cost=11.0), watched=True)                        # would need a huge move
        _, excluded, _ = self.build()
        self.assertEqual(excluded["B0NOSOURCE"], "No live EU source")
        self.assertEqual(excluded["B0NOSALES0"], "No sales evidence")
        self.assertIn("GATED", excluded["B0GATEDREC"])
        self.assertEqual(excluded["B0RESTRICT"], "Gated for us on Amazon")
        self.assertIn("Plug", excluded["B0PLUGITEM"])
        self.assertIn("older than", excluded["B0STALEREC"])
        self.assertEqual(excluded["B0ALREADYA"], "Already a lead today")
        self.assertEqual(excluded["B0TOOFAR00"], "Out of reach (needs more than a 40% move either way)")

    def test_a_product_a_hair_under_the_bar_is_reported_not_tracked(self):
        # ROI a fraction under 25%: it needs well under 2% either way, so an alert would be pure noise.
        edge = eu_cost_for_roi(flat_model(), 32.0)                               # the cost that gives exactly 25%
        self.add(record("B0HAIRUNDER", price=32.0, cost=round(edge * 1.004, 2), roi_90d=0), watched=True)   # ~24.9% ROI
        self.add(record("B0WELLUNDER", price=32.0, cost=17.4, roi_90d=0), watched=True)                   # ~19%: a real near miss
        rows, excluded, result = self.build()
        self.assertIn("B0HAIRUNDER", excluded)
        self.assertIn("Within 2% of being a lead", excluded["B0HAIRUNDER"])
        self.assertIn("B0WELLUNDER", rows)
        entry = next(e for e in result["excluded"] if e["asin"] == "B0HAIRUNDER")
        self.assertGreater(entry["roi_now"], 24.0)                                   # the ROI travels with it to the spreadsheet

    def test_only_the_latest_record_per_asin_counts(self):
        self.add(record("B0TWICEAAA", price=32.0, cost=24.0, scanned_at=NOW - timedelta(days=20)), watched=True)
        self.add(record("B0TWICEAAA", price=36.0, cost=24.0, scanned_at=NOW - timedelta(days=1)))
        rows, _, result = self.build()
        self.assertEqual(rows["B0TWICEAAA"]["uk_now"], 36.0)
        self.assertEqual(sum(1 for r in result["rows"] if r["asin"] == "B0TWICEAAA"), 1)

    def test_a_stored_fee_record_that_cannot_reproduce_its_roi_is_left_out_not_guessed(self):
        bad = record("B0FEEDATA0", category_name="", referral_rate_used=0.0)
        bad.roi = 21.0
        self.add(bad, watched=True)
        _, excluded, _ = self.build()
        self.assertIn("thresholds would be a guess", excluded["B0FEEDATA0"])

    def test_summary_counts_add_up(self):
        self.add(record("B0ONEONE00"), watched=True)
        self.add(record("B0NOSOURCE", cost=0.0, roi=15.0), watched=True)
        _, _, result = self.build()
        s = result["summary"]
        self.assertEqual(s["included"] + s["excluded"], s["candidates"])
        self.assertEqual(sum(s["excluded_by_reason"].values()), s["excluded"])


class FileTests(_Db):
    def test_files_are_plain_asin_lists_plus_a_spreadsheet(self):
        for i, (price, cost) in enumerate(((32, 24), (32, 24.5), (45, 36), (25, 20))):
            self.add(record(f"B0FILET{i}XX", price=price, cost=cost), watched=True)
        result = KeepaWatchListService.build(now=NOW)
        with tempfile.TemporaryDirectory() as out:
            paths = KeepaWatchListService.write_files(result, out)
            self.assertIn("keepa_watch_lists.xlsx", paths)
            self.assertIn("ALL_ASINS_for_a_first_test_upload.txt", paths)

            all_asins = open(paths["ALL_ASINS_for_a_first_test_upload.txt"], encoding="utf-8").read().split("\n")
            self.assertEqual(all_asins[-1], "")                                   # trailing newline only
            self.assertEqual(sorted(a for a in all_asins if a), sorted(r["asin"] for r in result["rows"]))

            for name, path in paths.items():
                if name.startswith(("EU_drop_", "UK_rise_")):
                    lines = [l for l in open(path, encoding="utf-8").read().split("\n") if l]
                    self.assertEqual(len(lines), len(set(lines)))               # no duplicates
                    self.assertTrue(all(len(l) == 10 and l.startswith("B0") for l in lines))
                    bucket = int(name.split("_")[-1][:2])
                    key = "eu_alert_pct" if name.startswith("EU_") else "uk_alert_pct"
                    expected = sorted(r["asin"] for r in result["rows"] if r[key] == bucket)
                    self.assertEqual(sorted(lines), expected)

            wb = load_workbook(paths["keepa_watch_lists.xlsx"])
            self.assertEqual(wb.sheetnames, ["Read me", "Products", "Left out"])
            self.assertEqual(wb["Products"].max_row - 1, len(result["rows"]))
            headers = [c.value for c in wb["Products"][1]]
            for needed in ("ASIN", "EU drop needed %", "EU alert at %", "UK rise needed %", "UK alert at %"):
                self.assertIn(needed, headers)

    def test_empty_buckets_write_no_file_and_an_empty_result_still_works(self):
        with tempfile.TemporaryDirectory() as out:
            result = KeepaWatchListService.build(now=NOW)
            paths = KeepaWatchListService.write_files(result, out)
            self.assertEqual(list(paths), ["keepa_watch_lists.xlsx"])
            self.assertEqual(load_workbook(paths["keepa_watch_lists.xlsx"])["Products"].max_row, 1)


if __name__ == "__main__":
    unittest.main()
