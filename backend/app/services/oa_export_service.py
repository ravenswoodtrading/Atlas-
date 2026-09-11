"""
Builds the "OA to Investigate" Excel export (2026-09-07, Tamara: "I
want to be able to export the list of OA items to investigate via an
excel sheet"). Shared by the download button (app/routes/review_queue.py)
and the standalone CLI script (export_oa_investigate.py) so there is
exactly one place this row list/formatting logic lives.

Read-only against the live DB -- reuses SellerWatchService.
list_oa_worth_investigating (the exact same query the Review Queue's
own "OA to Investigate" view runs) and SourcingClassifier.
assess_certainty (the same certainty rating shown in the detail panel),
so the export always matches what's on screen. No writes, no Keepa
calls, no cost.
"""
import json
from io import BytesIO
from urllib.parse import quote_plus

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from app.services.seller_watch_service import SellerWatchService
from app.services.sourcing_classifier import SourcingClassifier

FONT_NAME = "Arial"
HEADER_FILL = "1F4E78"

HEADERS = [
    "Date Found", "Competitor", "ASIN", "Product Title", "Brand", "Category",
    "Amazon Price (GBP)", "Amazon Price 90d Avg (GBP)", "Target Buy Price (GBP)", "Breakeven Price (GBP)",
    "EAN", "Monthly Sales", "Sales Rank Drops (30d)", "Certainty",
    "Atlas Score", "Atlas Confidence", "Atlas Recommendation",
    "Price Trend (90d)", "Demand Improving", "Competition Improving",
    "Currently Buyable", "Google Search", "Amazon Listing",
]
COLUMN_WIDTHS = [16, 20, 13, 42, 16, 16, 12, 14, 14, 14, 15, 12, 12, 10, 10, 10, 16, 12, 14, 18, 10, 10, 10]


def _google_search_url(title: str, ean: str, asin: str) -> str:
    # Same query-building convention as competitors.py's own
    # _google_search_url_variants -- title+EAN is the most precise
    # search when an EAN exists, falling back to ASIN+title otherwise.
    if ean:
        return f"https://www.google.com/search?q={quote_plus(f'{title} {ean}'.strip())}"
    return f"https://www.google.com/search?q={quote_plus(f'{asin} {title}'.strip())}"


def _build_rows() -> list[dict]:
    items = SellerWatchService.list_oa_worth_investigating(limit=1000)

    rows = []
    for entry in items:
        listing, seller, record, oa_price_guide = (
            entry["listing"], entry["seller"], entry["record"], entry["oa_price_guide"],
        )

        reasoning = {}
        if listing.sourcing_reasoning_json:
            try:
                reasoning = json.loads(listing.sourcing_reasoning_json)
            except Exception:
                reasoning = {}

        brand_pattern = SellerWatchService.brand_sourcing_pattern(record.brand)
        certainty = SourcingClassifier.assess_certainty(
            record.category_name, listing.sourcing_tag, reasoning, brand_pattern=brand_pattern,
        )

        # Trend deltas (price/demand/competition movement over the last
        # 90 days) were already computed by TrendEngine at scan time and
        # saved into report_json -- reading them back here costs nothing
        # extra and gives a real signal beyond a single point-in-time
        # price, without a fresh Keepa call.
        trend = {}
        if record.report_json:
            try:
                trend = json.loads(record.report_json).get("trend") or {}
            except Exception:
                trend = {}

        rows.append({
            "date_found": listing.detected_at,
            "competitor": seller.nickname or seller.seller_id,
            "asin": listing.asin,
            "title": record.title,
            "brand": record.brand,
            "category": record.category_name,
            "amazon_price_gbp": record.buy_box_now or None,
            "amazon_price_90d_gbp": record.buy_box_90d or None,
            "target_buy_price_gbp": oa_price_guide.get("target"),
            "breakeven_price_gbp": oa_price_guide.get("breakeven"),
            "ean": record.ean,
            "monthly_sales": record.monthly_sales,
            "sales_rank_drops_30d": record.sales_drops_30d,
            "certainty": certainty.get("level"),
            "atlas_score": record.score,
            "atlas_confidence": record.confidence,
            "atlas_recommendation": record.recommendation,
            "price_trend_90d_pct": trend.get("price_change"),
            "demand_improving": "Yes" if trend.get("demand_improving") else "No" if trend else "",
            "competition_improving": "Yes" if trend.get("competition_improving") else "No" if trend else "",
            "currently_buyable": "Yes" if listing.currently_buyable else "No",
            "google_search_url": _google_search_url(record.title, record.ean, listing.asin),
            "amazon_url": f"https://www.amazon.co.uk/dp/{listing.asin}",
        })

    # Most recently found first -- matches the Review Queue's own default sort.
    rows.sort(key=lambda r: r["date_found"] or "", reverse=True)
    return rows


def build_oa_investigate_workbook() -> tuple[BytesIO, int]:
    """Returns (workbook_bytes, row_count)."""
    rows = _build_rows()

    wb = Workbook()
    ws = wb.active
    ws.title = "OA to Investigate"

    ws.append(HEADERS)
    for col_idx in range(1, len(HEADERS) + 1):
        cell = ws.cell(row=1, column=col_idx)
        cell.font = Font(name=FONT_NAME, bold=True, color="FFFFFF")
        cell.fill = PatternFill(start_color=HEADER_FILL, end_color=HEADER_FILL, fill_type="solid")
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    ws.freeze_panes = "A2"

    for r in rows:
        row_idx = ws.max_row + 1
        ws.cell(row=row_idx, column=1, value=r["date_found"].strftime("%Y-%m-%d %H:%M") if r["date_found"] else "")
        ws.cell(row=row_idx, column=2, value=r["competitor"])
        ws.cell(row=row_idx, column=3, value=r["asin"])
        ws.cell(row=row_idx, column=4, value=r["title"])
        ws.cell(row=row_idx, column=5, value=r["brand"])
        ws.cell(row=row_idx, column=6, value=r["category"])
        ws.cell(row=row_idx, column=7, value=round(r["amazon_price_gbp"], 2) if r["amazon_price_gbp"] else None)
        ws.cell(row=row_idx, column=8, value=round(r["amazon_price_90d_gbp"], 2) if r["amazon_price_90d_gbp"] else None)
        ws.cell(row=row_idx, column=9, value=round(r["target_buy_price_gbp"], 2) if r["target_buy_price_gbp"] else None)
        ws.cell(row=row_idx, column=10, value=round(r["breakeven_price_gbp"], 2) if r["breakeven_price_gbp"] else None)
        ws.cell(row=row_idx, column=11, value=r["ean"] or "")
        ws.cell(row=row_idx, column=12, value=r["monthly_sales"])
        ws.cell(row=row_idx, column=13, value=r["sales_rank_drops_30d"])
        ws.cell(row=row_idx, column=14, value=r["certainty"])
        ws.cell(row=row_idx, column=15, value=r["atlas_score"])
        ws.cell(row=row_idx, column=16, value=r["atlas_confidence"])
        ws.cell(row=row_idx, column=17, value=r["atlas_recommendation"])
        ws.cell(row=row_idx, column=18, value=round(r["price_trend_90d_pct"], 1) if r["price_trend_90d_pct"] is not None else None)
        ws.cell(row=row_idx, column=19, value=r["demand_improving"])
        ws.cell(row=row_idx, column=20, value=r["competition_improving"])
        ws.cell(row=row_idx, column=21, value=r["currently_buyable"])

        gs_cell = ws.cell(row=row_idx, column=22, value="Search")
        gs_cell.hyperlink = r["google_search_url"]
        gs_cell.font = Font(name=FONT_NAME, color="0563C1", underline="single")

        az_cell = ws.cell(row=row_idx, column=23, value="View")
        az_cell.hyperlink = r["amazon_url"]
        az_cell.font = Font(name=FONT_NAME, color="0563C1", underline="single")

    for row in ws.iter_rows(min_row=2, max_row=ws.max_row, min_col=1, max_col=21):
        for cell in row:
            cell.font = Font(name=FONT_NAME)

    for idx, width in enumerate(COLUMN_WIDTHS, start=1):
        ws.column_dimensions[get_column_letter(idx)].width = width

    for price_col in (7, 8, 9, 10):
        for row in ws.iter_rows(min_row=2, max_row=ws.max_row, min_col=price_col, max_col=price_col):
            for cell in row:
                cell.number_format = '"£"#,##0.00'

    for row in ws.iter_rows(min_row=2, max_row=ws.max_row, min_col=18, max_col=18):
        for cell in row:
            cell.number_format = '0.0"%"'

    buf = BytesIO()
    wb.save(buf)
    buf.seek(0)
    return buf, len(rows)
