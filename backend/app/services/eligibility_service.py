"""
Eligibility Check (2026-09-23) -- "can this account list these ASINs?" for
an uploaded Keepa Product Finder export, or for a list of barcodes that
first have to be resolved to ASINs.

No eligibility logic of its own: every answer comes from
RestrictionService.check (same cache, same TTLs, same SP-API call as the
Scan Queue's pre-Keepa filter and the Review Queue's gated-ASIN guard), so
this page can never disagree with what the rest of Atlas believes.

Output columns are a contract with the oa-va Claude skill -- don't rename
them or change the values without checking that side:
  Eligible            Y (can list) / N (restricted) / Unknown (couldn't check).
                      Unknown is never folded into Y or N -- same rule as
                      RestrictionService's None.
  Restriction reason  reason_code as stored ("APPROVAL_REQUIRED+NOT_ELIGIBLE"),
                      blank unless Eligible is N.
  Checked at          ISO-8601 UTC of the answer ("2026-09-23T10:15:00Z") -- a
                      cache hit shows when it was really checked, not now.

Checks run as a background job (a cold 5,000-row Keepa export is ~40
minutes of 0.5s calls, far past any HTTP timeout); the page polls
job_status(). Jobs live in memory only -- a restart drops them, and the
answers themselves are already in ListingRestriction.
"""
import csv
import io
import re
import threading
import uuid
from collections import Counter
from datetime import datetime, timezone

from openpyxl import load_workbook

from app.database.database import SessionLocal
from app.database.models import ProductRecord
from app.services.product_repository import ProductRepository
from app.services.restriction_service import RestrictionService
from app.sp_api.client import get_sp_api_client

ELIGIBLE_COL = "Eligible"
REASON_COL = "Restriction reason"
CHECKED_COL = "Checked at"
OUTPUT_COLUMNS = (ELIGIBLE_COL, REASON_COL, CHECKED_COL)

ELIGIBLE, RESTRICTED, UNKNOWN = "Y", "N", "Unknown"

# Brands with at least this many restricted ASINs in one file are listed
# as candidates for the gated-brands list.
NOTABLE_BRAND_MIN_RESTRICTED = 3

# RestrictionService.check is called in chunks this size so the page can
# show progress; the cache/API behaviour is identical to one big call.
CHECK_CHUNK_SIZE = 25

MAX_JOBS_KEPT = 20
MAX_BARCODES = 5000

_ASIN_RE = re.compile(r"^[A-Z0-9]{10}$")
_ASIN_HEADERS = ("asin",)
_BRAND_HEADERS = ("brand",)
_BARCODE_HEADERS = ("ean", "upc", "gtin", "barcode", "product codes: ean", "product codes: upc",
                    "product codes: gtin", "ean13", "ean-13")


# --------------------------------------------------------------- answers

def _normalise_header(value) -> str:
    return " ".join(str(value or "").split()).lower()


def find_column(headers: list, names: tuple):
    """Index of the first header matching any of `names` (case- and
    whitespace-insensitive, exact), else None. Exact on purpose: Keepa
    exports also carry "Parent ASIN" and "Brand Store Name"."""
    normalised = [_normalise_header(h) for h in headers]
    for name in names:
        if name in normalised:
            return normalised.index(name)
    return None


def _iso(value) -> str:
    return value.strftime("%Y-%m-%dT%H:%M:%SZ") if value else ""


def eligibility_columns(asin: str, results: dict, details: dict) -> tuple:
    """(Eligible, Restriction reason, Checked at) for one ASIN, from
    RestrictionService.check's {asin: True/False/None} and .details()."""
    status = results.get(asin)
    if status is None:
        return UNKNOWN, "", ""

    detail = details.get(asin) or {}
    checked_at = _iso(detail.get("checked_at"))
    if status is True:
        return RESTRICTED, detail.get("reason_code", ""), checked_at
    return ELIGIBLE, "", checked_at


def check_asins(asins: list, marketplace: str = "UK", progress=None) -> tuple:
    """Returns (results, details) for every valid ASIN. `progress(done)` is
    called after each chunk."""
    unique = list(dict.fromkeys(a for a in asins if a))
    results = {}

    for i in range(0, len(unique), CHECK_CHUNK_SIZE):
        results.update(RestrictionService.check(unique[i:i + CHECK_CHUNK_SIZE], marketplace))
        if progress:
            progress(min(i + CHECK_CHUNK_SIZE, len(unique)))

    return results, RestrictionService.details(unique, marketplace)


def notable_brands(pairs: list, minimum: int = NOTABLE_BRAND_MIN_RESTRICTED) -> list:
    """pairs: [(asin, brand, eligible)] -- one per distinct ASIN. Brands with
    at least `minimum` restricted ASINs, most restricted first. A blank
    brand is filled from ProductRecord where Atlas has already scanned the
    ASIN."""
    missing = [asin for asin, brand, _ in pairs if not (brand or "").strip()]
    known = {}
    if missing:
        db = SessionLocal()
        try:
            for i in range(0, len(missing), 500):
                for asin, brand in (db.query(ProductRecord.asin, ProductRecord.brand)
                                    .filter(ProductRecord.asin.in_(missing[i:i + 500]), ProductRecord.brand != "")
                                    .all()):
                    known.setdefault(asin, brand)
        finally:
            db.close()

    restricted, checked, display = Counter(), Counter(), {}
    for asin, brand, eligible in pairs:
        brand = ((brand or "").strip() or known.get(asin, "")).strip()
        if not brand or eligible == UNKNOWN:
            continue
        key = brand.lower()
        display.setdefault(key, brand)
        checked[key] += 1
        if eligible == RESTRICTED:
            restricted[key] += 1

    try:
        gated = {brand for brand, category_id in ProductRepository.get_gated_brand_pairs() if not category_id}
    except Exception as exc:
        print(f"[eligibility] gated brand lookup failed (non-fatal): {exc}")
        gated = set()

    rows = [dict(brand=display[key], restricted=count, checked=checked[key], already_gated=key in gated)
            for key, count in restricted.items() if count >= minimum]
    return sorted(rows, key=lambda r: (-r["restricted"], r["brand"].lower()))


def _tally(values: list) -> dict:
    counts = Counter(values)
    return {ELIGIBLE: counts[ELIGIBLE], RESTRICTED: counts[RESTRICTED], UNKNOWN: counts[UNKNOWN]}


# ----------------------------------------------------------- file in/out

def _decode(data: bytes) -> str:
    for encoding in ("utf-8-sig", "cp1252"):
        try:
            return data.decode(encoding)
        except UnicodeDecodeError:
            continue
    return data.decode("latin-1")


def _is_xlsx(filename: str) -> bool:
    return (filename or "").lower().endswith((".xlsx", ".xlsm"))


def _csv_rows(data: bytes) -> tuple:
    """(rows, dialect) -- the dialect is reused to write the file back."""
    text = _decode(data)
    try:
        dialect = csv.Sniffer().sniff(text[:4096], delimiters=",;\t")
    except csv.Error:
        dialect = csv.excel
    return list(csv.reader(io.StringIO(text), dialect)), dialect


def read_table(filename: str, data: bytes) -> tuple:
    """(headers, rows) as lists of cell values, first sheet for XLSX."""
    if _is_xlsx(filename):
        ws = load_workbook(io.BytesIO(data), read_only=True, data_only=True).active
        rows = [list(r) for r in ws.iter_rows(values_only=True)]
    else:
        rows, _ = _csv_rows(data)
    if not rows:
        raise ValueError("The file is empty.")
    return list(rows[0]), rows[1:]


def _output_column_indexes(headers: list) -> list:
    """Where the three output columns go: over existing ones (re-uploading
    an enriched file), else appended after the last header."""
    indexes, next_free = [], len(headers)
    for name in OUTPUT_COLUMNS:
        existing = find_column(headers, (name.lower(),))
        if existing is None:
            existing, next_free = next_free, next_free + 1
        indexes.append(existing)
    return indexes


def _enrich_csv(data: bytes, values: list) -> bytes:
    rows, dialect = _csv_rows(data)
    headers = rows[0]
    indexes = _output_column_indexes(headers)
    width = max(indexes) + 1

    out = io.StringIO()
    writer = csv.writer(out, dialect, lineterminator="\r\n")
    header_row = headers + [""] * (width - len(headers))
    for idx, name in zip(indexes, OUTPUT_COLUMNS):
        header_row[idx] = name
    writer.writerow(header_row)
    for row, triple in zip(rows[1:], values):
        row = row + [""] * (width - len(row))
        for idx, value in zip(indexes, triple):
            row[idx] = value
        writer.writerow(row)
    return out.getvalue().encode("utf-8-sig")


def _enrich_xlsx(data: bytes, values: list) -> bytes:
    wb = load_workbook(io.BytesIO(data))
    ws = wb.active
    headers = [c.value for c in ws[1]]
    indexes = _output_column_indexes(headers)
    for idx, name in zip(indexes, OUTPUT_COLUMNS):
        ws.cell(row=1, column=idx + 1, value=name)
    for row_idx, triple in enumerate(values, start=2):
        for idx, value in zip(indexes, triple):
            ws.cell(row=row_idx, column=idx + 1, value=value or None)
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def _clean_asin(value) -> str:
    return str(value or "").strip().upper()


def enrich_keepa_export(filename: str, data: bytes, marketplace: str = "UK", progress=None) -> dict:
    """The whole Part 1 pipeline. Returns {"content", "summary"}; raises
    ValueError for a file it can't use."""
    headers, rows = read_table(filename, data)
    asin_col = find_column(headers, _ASIN_HEADERS)
    if asin_col is None:
        raise ValueError("No ASIN column found -- the header row needs a column called 'ASIN'.")
    brand_col = find_column(headers, _BRAND_HEADERS)

    row_asins = [_clean_asin(r[asin_col]) if asin_col < len(r) else "" for r in rows]
    valid = [a for a in row_asins if _ASIN_RE.match(a)]
    if progress:
        progress(0, len(set(valid)))

    results, details = check_asins(valid, marketplace, progress=(lambda done: progress(done, None)) if progress else None)

    values = []
    for asin in row_asins:
        if not asin:
            values.append(("", "", ""))           # not a product row -- nothing to check
        elif not _ASIN_RE.match(asin):
            values.append((UNKNOWN, "", ""))      # not a checkable ASIN
        else:
            values.append(eligibility_columns(asin, results, details))

    content = (_enrich_xlsx(data, values) if _is_xlsx(filename)
               else _enrich_csv(data, values))

    seen, pairs = set(), []
    for row, asin, triple in zip(rows, row_asins, values):
        if _ASIN_RE.match(asin) and asin not in seen:
            seen.add(asin)
            brand = row[brand_col] if brand_col is not None and brand_col < len(row) else ""
            pairs.append((asin, str(brand or ""), triple[0]))

    return dict(content=content, summary=dict(
        rows=len(rows), asins=len(pairs), counts=_tally([p[2] for p in pairs]),
        invalid_rows=sum(1 for a in row_asins if a and not _ASIN_RE.match(a)),
        brand_column=headers[brand_col] if brand_col is not None else None,
        notable_brands=notable_brands(pairs),
    ))


# ------------------------------------------------------- barcode -> ASIN

def normalise_barcode(raw) -> tuple:
    """(code as entered, EAN-13/EAN-8 to query) -- query is "" if the code
    isn't a usable barcode. UPC-12 is sent as its EAN-13 (leading 0):
    verified live 2026-09-23 that both forms resolve to the same ASIN."""
    if isinstance(raw, float) and raw.is_integer():
        raw = int(raw)                              # Excel numeric cell
    entered = str(raw or "").strip()
    digits = re.sub(r"\D", "", entered)
    if len(digits) == 12:
        return entered, "0" + digits
    if len(digits) == 14 and digits.startswith("0"):
        return entered, digits[1:]
    if len(digits) in (8, 13):
        return entered, digits
    return entered, ""


def parse_barcodes(text: str = "", filename: str = "", data: bytes = b"") -> list:
    """Codes in entry order, de-duplicated: from pasted text (any
    whitespace/comma/semicolon separated) and/or a CSV/XLSX with an
    EAN/UPC/GTIN/Barcode column (or a single column with no header)."""
    codes = [c for c in re.split(r"[\s,;]+", text or "") if c]

    if data:
        headers, rows = read_table(filename, data)
        col = find_column(headers, _BARCODE_HEADERS)
        if col is None:
            if len(headers) == 1 and normalise_barcode(headers[0])[1]:
                col, rows = 0, [headers] + rows
            else:
                raise ValueError("No barcode column found -- name it EAN, UPC, GTIN or Barcode.")
        for row in rows:
            cell = row[col] if col < len(row) else None
            if isinstance(cell, float):
                cell = normalise_barcode(cell)[0]
            # Keepa's "Product Codes: EAN" can hold several, comma-separated.
            codes.extend(c for c in re.split(r"[\s,;]+", str(cell or "")) if c)

    return list(dict.fromkeys(codes))


BARCODE_HEADERS_OUT = ["Barcode", "Match count", "Needs check", "ASIN", "Title", "Brand",
                       ELIGIBLE_COL, REASON_COL, CHECKED_COL]


def resolve_barcodes(codes: list, marketplace: str = "UK", sp_client=None, progress=None) -> dict:
    """The whole Part 2 pipeline: barcode -> ASIN(s) via Catalog Items, then
    the same RestrictionService path as a Keepa export. Returns
    {"content" (CSV bytes), "summary"}.

    One output row per (barcode, ASIN) match. A barcode that matched
    nothing or couldn't be looked up gets one row with no ASIN and blank
    eligibility (there's nothing to check); Needs check = Y for anything
    other than exactly one match."""
    client = sp_client or get_sp_api_client()
    if client is None:
        raise ValueError("SP-API isn't configured, so barcodes can't be looked up.")

    entries = [normalise_barcode(c) for c in codes[:MAX_BARCODES]]
    queries = list(dict.fromkeys(q for _, q in entries if q))
    batch = client.CATALOG_ITEMS_BATCH_SIZE
    total_steps = max(1, -(-len(queries) // batch))
    if progress:
        progress(0, total_steps)

    matches = {}            # query -> list, or None if the lookup failed
    for i in range(0, len(queries), batch):
        chunk = queries[i:i + batch]
        answer = client.search_catalog_items_by_identifier(chunk, "EAN", marketplace)
        for q in chunk:
            matches[q] = None if answer is None else answer.get(q, [])
        if progress:
            progress(i // batch + 1, None)

    asins = list(dict.fromkeys(m["asin"] for found in matches.values() if found for m in found))
    if progress:
        progress(0, len(asins))
    results, details = check_asins(asins, marketplace, progress=(lambda done: progress(done, None)) if progress else None)

    out = io.StringIO()
    writer = csv.writer(out, lineterminator="\r\n")
    writer.writerow(BARCODE_HEADERS_OUT)
    flags, pairs, seen = [], [], set()
    status_counts = Counter()

    for entered, query in entries:
        found = matches.get(query) if query else None
        if not query:
            status, found = "Invalid barcode", []
        elif found is None:
            status, found = "Lookup failed", []
        else:
            status = {0: "No match", 1: "OK"}.get(len(found), "Multiple matches")
        status_counts[status] += 1
        if status != "OK":
            flags.append(dict(barcode=entered, status=status, asins=[m["asin"] for m in found]))

        if not found:
            count = "" if status in ("Invalid barcode", "Lookup failed") else 0
            writer.writerow([entered, count, "Y", "", status, "", "", "", ""])
            continue
        for m in found:
            triple = eligibility_columns(m["asin"], results, details)
            writer.writerow([entered, len(found), "Y" if len(found) != 1 else "", m["asin"],
                             m["title"], m["brand"], *triple])
            if m["asin"] not in seen:
                seen.add(m["asin"])
                pairs.append((m["asin"], m["brand"], triple[0]))

    return dict(content=out.getvalue().encode("utf-8-sig"), summary=dict(
        barcodes=len(entries), truncated=max(0, len(codes) - MAX_BARCODES), asins=len(pairs),
        counts=_tally([p[2] for p in pairs]), statuses=dict(status_counts), flags=flags,
        notable_brands=notable_brands(pairs),
    ))


# ------------------------------------------------------------------ jobs

_jobs = {}
_jobs_lock = threading.Lock()
# One check at a time: every job shares the one SP-API client and its
# pacing, so two at once would only each run at half speed.
_run_lock = threading.Lock()


def _new_job(kind: str, label: str, out_name: str, media_type: str) -> dict:
    job = dict(id=uuid.uuid4().hex[:12], kind=kind, label=label, out_name=out_name, media_type=media_type,
               status="queued", stage="Waiting for another check to finish", done=0, total=0,
               created_at=datetime.now(timezone.utc), finished_at=None, error="", summary=None, content=None)
    with _jobs_lock:
        _jobs[job["id"]] = job
        for old in sorted(_jobs.values(), key=lambda j: j["created_at"])[:-MAX_JOBS_KEPT]:
            _jobs.pop(old["id"], None)
    return job


def _run(job: dict, work, stages: list):
    def progress(done, total):
        if total is not None:
            job["stage"] = stages.pop(0) if stages else job["stage"]
            job["total"], job["done"] = total, 0
        else:
            job["done"] = done

    def target():
        with _run_lock:
            job["status"], job["stage"] = "running", stages[0] if stages else "Checking"
            try:
                result = work(progress)
                job["content"], job["summary"], job["status"] = result["content"], result["summary"], "done"
            except ValueError as exc:
                job["status"], job["error"] = "error", str(exc)
            except Exception as exc:
                print(f"[eligibility] job {job['id']} failed: {exc!r}")
                job["status"], job["error"] = "error", f"Unexpected error: {exc}"
            finally:
                job["finished_at"] = datetime.now(timezone.utc)

    threading.Thread(target=target, name=f"eligibility-{job['id']}", daemon=True).start()


def start_keepa_job(filename: str, data: bytes) -> dict:
    read_table(filename, data)   # fail fast on an unreadable file, before queueing
    base = re.sub(r"[^A-Za-z0-9._() -]", "_", (filename or "keepa_export").rsplit(".", 1)[0]) or "keepa_export"
    ext = "xlsx" if _is_xlsx(filename) else "csv"
    media = ("application/vnd.openxmlformats-officedocument.spreadsheetml.sheet" if ext == "xlsx"
             else "text/csv")
    job = _new_job("keepa", filename or "Keepa export", f"{base}_eligibility.{ext}", media)
    _run(job, lambda progress: enrich_keepa_export(filename, data, progress=progress),
         ["Checking ASINs"])
    return job


def start_barcode_job(codes: list) -> dict:
    if not codes:
        raise ValueError("No barcodes found.")
    stamp = datetime.now(timezone.utc).strftime("%Y%m%d_%H%M")
    job = _new_job("barcodes", f"{len(codes)} barcode(s)", f"barcode_eligibility_{stamp}.csv", "text/csv")
    _run(job, lambda progress: resolve_barcodes(codes, progress=progress),
         ["Looking up barcodes on Amazon", "Checking ASINs"])
    return job


def get_job(job_id: str):
    with _jobs_lock:
        return _jobs.get(job_id)


def recent_jobs() -> list:
    with _jobs_lock:
        return sorted(_jobs.values(), key=lambda j: j["created_at"], reverse=True)
