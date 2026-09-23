"""Eligibility Check: Keepa export enrichment and barcode -> ASIN resolution.
Builds on test_listing_restrictions' isolated SQLite and fake SP-API client --
no live database writes or API calls."""
import csv
import io
import unittest
from unittest.mock import MagicMock, patch

from openpyxl import Workbook, load_workbook

from app.database.models import ProductRecord
from app.services import eligibility_service as es
from app.services import restriction_service as rs
from app.sp_api import client as client_module
from app.sp_api.client import SPAPIClient
from test_listing_restrictions import _DbCase, _response

OK, GATED, FAILS = "B000000OK1", "B00GATED01", "B000FAILS1"


def _csv(rows):
    out = io.StringIO()
    csv.writer(out).writerows(rows)
    return out.getvalue().encode("utf-8-sig")


def _read_csv(data):
    return list(csv.reader(io.StringIO(data.decode("utf-8-sig"))))


class _EligibilityCase(_DbCase):
    def setUp(self):
        super().setUp()
        patcher = patch.object(es, "SessionLocal", self.sessions)
        patcher.start()
        self.addCleanup(patcher.stop)
        self.gated = set()
        patcher = patch.object(es.ProductRepository, "get_gated_brand_pairs", side_effect=lambda: self.gated)
        patcher.start()
        self.addCleanup(patcher.stop)

    def use_sp(self, answers):
        sp = self.sp(answers)
        patcher = patch.object(rs, "get_sp_api_client", return_value=sp)
        patcher.start()
        self.addCleanup(patcher.stop)
        return sp


class EligibleValueTests(_EligibilityCase):
    def test_three_values_never_collapse_unknown(self):
        self.use_sp({OK: False, GATED: True, FAILS: None})
        results, details = es.check_asins([OK, GATED, FAILS])
        self.assertEqual(es.eligibility_columns(OK, results, details)[:2], ("Y", ""))
        self.assertEqual(es.eligibility_columns(GATED, results, details)[:2], ("N", "X"))
        self.assertEqual(es.eligibility_columns(FAILS, results, details), ("Unknown", "", ""))

    def test_cache_hit_reports_when_it_was_really_checked(self):
        self.store(GATED, True, age_days=3)
        sp = self.use_sp({})
        results, details = es.check_asins([GATED])
        eligible, reason, checked_at = es.eligibility_columns(GATED, results, details)
        self.assertEqual((eligible, reason), ("N", "APPROVAL_REQUIRED"))
        expected = details[GATED]["checked_at"].strftime("%Y-%m-%dT%H:%M:%SZ")
        self.assertEqual(checked_at, expected)
        self.assertLess(details[GATED]["checked_at"], rs._now() - rs.timedelta(days=2))
        sp.get_listing_restrictions.assert_not_called()

    def test_details_follow_the_same_ttl_as_cached(self):
        self.store("R_OLD", True, age_days=8)
        self.store("OK_FRESH", False, age_days=13)
        self.assertEqual(set(rs.RestrictionService.details(["R_OLD", "OK_FRESH"])), {"OK_FRESH"})


class KeepaExportTests(_EligibilityCase):
    HEADERS = ["Locale", "Title", "asin ", "Parent ASIN", "New\n 3rd Party FBA: Current", "Brand", "Brand Store Name"]

    def rows(self):
        return [self.HEADERS,
                ["uk", "Thing, with comma", OK, "B0PARENT01", "9.99", "Acme", "Acme Store"],
                ["uk", "Gated", GATED, "", "", "Acme", ""],
                ["uk", "Failed", FAILS, "", "", "", ""],
                ["uk", "Duplicate", OK.lower(), "", "", "Acme", ""],
                ["uk", "Bad asin", "nonsense", "", "", "", ""],
                ["uk", "Blank row", "", "", "", "", ""]]

    def test_csv_drops_restricted_rows_and_keeps_everything_else_untouched(self):
        sp = self.use_sp({OK: False, GATED: True, FAILS: None})
        result = es.enrich_keepa_export("export.csv", _csv(self.rows()))
        out = _read_csv(result["content"])
        self.assertEqual(out[0], self.HEADERS + ["Eligible", "Restriction reason", "Checked at"])
        kept = [r for r in self.rows()[1:] if r[2] != GATED]
        self.assertEqual([r[:7] for r in out[1:]], kept)                 # untouched, same order, GATED gone
        self.assertEqual([r[7] for r in out[1:]], ["Y", "Unknown", "Y", "Unknown", ""])
        self.assertEqual([r[8] for r in out[1:]], ["", "", "", "", ""])
        self.assertTrue(out[1][9].endswith("Z") and out[2][9] == "")
        self.assertEqual(sorted(c.args[0] for c in sp.get_listing_restrictions.call_args_list), sorted([OK, GATED, FAILS]))
        self.assertEqual(result["summary"]["counts"], {"Y": 1, "N": 1, "Unknown": 1})   # removed rows still counted
        self.assertEqual(result["summary"]["removed_rows"], 1)
        self.assertEqual(result["summary"]["invalid_rows"], 1)

    def test_every_row_of_a_restricted_asin_is_dropped(self):
        self.use_sp({OK: False, GATED: True})
        rows = [["ASIN"], [GATED], [OK], [GATED.lower()]]
        result = es.enrich_keepa_export("x.csv", _csv(rows))
        self.assertEqual([r[0] for r in _read_csv(result["content"])[1:]], [OK])
        self.assertEqual(result["summary"]["removed_rows"], 2)

    def test_reuploading_an_enriched_file_overwrites_instead_of_duplicating(self):
        self.use_sp({OK: False})
        first = [["ASIN", "Eligible", "restriction reason", "Checked at", "Notes"], [OK, "N", "OLD", "old", "keep"]]
        out = _read_csv(es.enrich_keepa_export("x.csv", _csv(first))["content"])
        self.assertEqual(out[0], ["ASIN", "Eligible", "Restriction reason", "Checked at", "Notes"])  # contract casing
        self.assertEqual(out[1][:3] + out[1][4:], [OK, "Y", "", "keep"])
        self.assertTrue(out[1][3].endswith("Z"))

    def test_no_asin_column_is_a_clear_error(self):
        with self.assertRaisesRegex(ValueError, "No ASIN column"):
            es.enrich_keepa_export("x.csv", _csv([["Parent ASIN", "Title"], ["B0PARENT01", "x"]]))

    def test_xlsx_drops_restricted_rows_and_slides_the_rest_up_with_formatting(self):
        from openpyxl.styles import Font
        ok2 = "B000000OK2"
        self.use_sp({OK: False, GATED: True, ok2: None})
        wb = Workbook()
        ws = wb.active
        ws.append(["Title", "ASIN", "Price"])
        ws.append(["a", OK, 1.5])
        ws.append(["b", GATED, 2.5])
        ws.append(["c", ok2, 3.5])
        ws["C4"].number_format = '"£"#,##0.00'
        ws["A4"].font = Font(bold=True)
        ws["A4"].hyperlink = "https://www.amazon.co.uk/dp/B000000OK2"
        buf = io.BytesIO()
        wb.save(buf)
        out_ws = load_workbook(io.BytesIO(es.enrich_keepa_export("export.xlsx", buf.getvalue())["content"])).active
        rows = list(out_ws.iter_rows(values_only=True))
        self.assertEqual(rows[0], ("Title", "ASIN", "Price", "Eligible", "Restriction reason", "Checked at"))
        self.assertEqual(rows[1][:5], ("a", OK, 1.5, "Y", None))
        self.assertEqual(rows[2][:5], ("c", ok2, 3.5, "Unknown", None))    # moved up from row 4
        self.assertEqual(len(rows), 3)                                     # nothing left behind
        self.assertEqual(out_ws["C3"].number_format, '"£"#,##0.00')
        self.assertTrue(out_ws["A3"].font.bold)
        self.assertEqual(out_ws["A3"].hyperlink.target, "https://www.amazon.co.uk/dp/B000000OK2")
        self.assertIsNone(out_ws["A2"].hyperlink)

    def test_quote_marks_after_the_sniffed_sample_still_read_and_write(self):
        # Real Keepa export (2026-09-23): no "" in the first 4KB made the Sniffer
        # pick doublequote=False, then a later 24" title failed to write back.
        self.use_sp({OK: False})
        rows = [["Title", "ASIN"]] + [[f"Plain title number {i} " * 3, OK] for i in range(80)]
        rows += [['PHILIPS 24" FHD Monitor, "White"', OK]]
        out = _read_csv(es.enrich_keepa_export("x.csv", _csv(rows))["content"])
        self.assertEqual(out[-1][:3], ['PHILIPS 24" FHD Monitor, "White"', OK, "Y"])
        self.assertEqual(len(out), len(rows))

    def test_semicolon_csv_is_read_and_written_back_the_same_way(self):
        self.use_sp({OK: False})
        data = f"Title;ASIN\r\nx;{OK}\r\n".encode("utf-8")
        text = es.enrich_keepa_export("x.csv", data)["content"].decode("utf-8-sig")
        self.assertTrue(text.startswith("Title;ASIN;Eligible;Restriction reason;Checked at"))


class ProgressAndSavedFileTests(_EligibilityCase):
    def test_progress_reports_how_many_need_a_real_call(self):
        self.store(OK, False)
        self.use_sp({GATED: True, FAILS: None})
        calls = []
        es.check_asins([OK, GATED, FAILS], progress=lambda done, total, to_ask=None: calls.append((done, total, to_ask)))
        self.assertEqual(calls[0], (0, 3, 2))
        self.assertEqual(calls[-1], (3, None, None))

    def test_finished_files_are_saved_listed_pruned_and_served_by_exact_name_only(self):
        import tempfile
        from pathlib import Path
        with tempfile.TemporaryDirectory() as tmp, patch.object(es, "EXPORT_DIR", Path(tmp) / "elig"), \
                patch.object(es, "MAX_SAVED_FILES", 2):
            for i, job_id in enumerate(("aaa", "bbb", "ccc")):
                es._save_output(dict(id=job_id, out_name=f"export{i}_eligibility.csv", content=b"x"))
                path = es.EXPORT_DIR / f"{job_id}__export{i}_eligibility.csv"
                import os
                os.utime(path, (1000 + i, 1000 + i))
            es._save_output(dict(id="ddd", out_name="new.csv", content=b"y"))
            names = [f["name"] for f in es.saved_files()]
            self.assertEqual(len(names), 2)
            self.assertEqual(names[0], "ddd__new.csv")
            self.assertEqual(es.saved_files()[0]["label"], "new.csv")
            self.assertIsNotNone(es.saved_file_path("ddd__new.csv"))
            self.assertIsNone(es.saved_file_path("../ddd__new.csv"))
            self.assertIsNone(es.saved_file_path("aaa__export0_eligibility.csv"))   # pruned


class NotableBrandTests(_EligibilityCase):
    def test_brands_with_several_restricted_asins_are_listed_with_gated_flag(self):
        pairs = [(f"A{i}", "Acme", "N") for i in range(3)] + [("A9", "acme", "Y")]
        pairs += [(f"B{i}", "Bolt", "N") for i in range(4)] + [("C0", "Cog", "N"), ("C1", "Cog", "N")]
        pairs += [("U0", "Acme", "Unknown")]
        self.gated = {("bolt", ""), ("cog", "cat123")}
        rows = es.notable_brands(pairs)
        self.assertEqual(rows, [
            dict(brand="Bolt", restricted=4, checked=4, already_gated=True),
            dict(brand="Acme", restricted=3, checked=4, already_gated=False),
        ])

    def test_blank_brand_falls_back_to_product_record(self):
        with self.sessions() as db:
            for asin in ("P1", "P2", "P3"):
                db.add(ProductRecord(asin=asin, brand="Philips"))
            db.commit()
        rows = es.notable_brands([("P1", "", "N"), ("P2", " ", "N"), ("P3", "", "N")])
        self.assertEqual([(r["brand"], r["restricted"]) for r in rows], [("Philips", 3)])


class BarcodeParsingTests(unittest.TestCase):
    def test_normalise(self):
        self.assertEqual(es.normalise_barcode("8719514459472"), ("8719514459472", "8719514459472"))
        self.assertEqual(es.normalise_barcode("019659200341"), ("019659200341", "0019659200341"))
        self.assertEqual(es.normalise_barcode("00012345678905"), ("00012345678905", "0012345678905"))
        self.assertEqual(es.normalise_barcode(8719514459472.0), ("8719514459472", "8719514459472"))
        self.assertEqual(es.normalise_barcode("12345")[1], "")

    def test_pasted_text_and_keepa_style_column(self):
        data = _csv([["ASIN", "Product Codes: EAN"], ["B0X", "111,222"], ["B0Y", ""], ["B0Z", "333"]])
        codes = es.parse_barcodes("999 111\n444;", "k.csv", data)
        self.assertEqual(codes, ["999", "111", "444", "222", "333"])

    def test_headerless_single_column(self):
        self.assertEqual(es.parse_barcodes(filename="c.csv", data=_csv([["8719514459472"], ["5000000000001"]])),
                         ["8719514459472", "5000000000001"])


class ResolveBarcodeTests(_EligibilityCase):
    def fake_catalog(self, table, fail_codes=()):
        sp = MagicMock()
        sp.CATALOG_ITEMS_BATCH_SIZE = 1

        def lookup(codes, identifiers_type, marketplace):
            if any(c in fail_codes for c in codes):
                return None
            return {c: table.get(c, []) for c in codes}

        sp.search_catalog_items_by_identifier.side_effect = lookup
        return sp

    def test_rows_flags_and_same_eligibility_path(self):
        self.use_sp({"B0ONE00001": False, "B0MULTI001": True, "B0MULTI002": None})
        catalog = self.fake_catalog({
            "5000000000001": [dict(asin="B0ONE00001", title="One", brand="Acme")],
            "5000000000002": [dict(asin="B0MULTI001", title="M1", brand="Acme"),
                              dict(asin="B0MULTI002", title="M2", brand="Acme")],
        }, fail_codes={"5000000000004"})
        codes = ["5000000000001", "5000000000002", "5000000000003", "5000000000004", "123"]
        result = es.resolve_barcodes(codes, sp_client=catalog)
        rows = _read_csv(result["content"])
        self.assertEqual(rows[0], es.BARCODE_HEADERS_OUT)
        body = [(r[0], r[1], r[2], r[3], r[6], r[7]) for r in rows[1:]]
        self.assertEqual(body, [
            ("5000000000001", "1", "", "B0ONE00001", "Y", ""),
            # B0MULTI001 is restricted, so its row is left out of the file
            ("5000000000002", "2", "Y", "B0MULTI002", "Unknown", ""),
            ("5000000000003", "0", "Y", "", "", ""),
            ("5000000000004", "", "Y", "", "", ""),    # lookup failed: not "no match"
            ("123", "", "Y", "", "", ""),
        ])
        self.assertEqual([r[4] for r in rows[3:]], ["No match", "Lookup failed", "Invalid barcode"])
        summary = result["summary"]
        self.assertEqual(summary["statuses"], {"OK": 1, "Multiple matches": 1, "No match": 1,
                                               "Lookup failed": 1, "Invalid barcode": 1})
        self.assertEqual(summary["counts"], {"Y": 1, "N": 1, "Unknown": 1})
        self.assertEqual(summary["removed_rows"], 1)
        self.assertEqual(len(summary["flags"]), 4)

    def test_no_sp_client_is_a_clear_error(self):
        with patch.object(es, "get_sp_api_client", return_value=None):
            with self.assertRaisesRegex(ValueError, "SP-API"):
                es.resolve_barcodes(["5000000000001"])


class CatalogIdentifierClientTests(unittest.TestCase):
    def setUp(self):
        self.client = SPAPIClient("id", "secret", "token", seller_id="SELLER1")
        for target, value in (("_get_access_token", "tok"), ("_pace", None)):
            patcher = patch.object(self.client, target, return_value=value)
            patcher.start()
            self.addCleanup(patcher.stop)
        for patcher in (patch.object(client_module.time, "sleep"),
                        patch.object(SPAPIClient, "_identifier_response_logged", True)):
            patcher.start()
            self.addCleanup(patcher.stop)

    @staticmethod
    def item(asin, *codes, brand="Philips"):
        return {"asin": asin,
                "identifiers": [{"marketplaceId": "A1F83G8C2ARO7P",
                                 "identifiers": [{"identifierType": "EAN", "identifier": c} for c in codes]}],
                "summaries": [{"marketplaceId": "A1F83G8C2ARO7P", "itemName": f"Title {asin}", "brand": brand}]}

    def test_items_are_mapped_back_by_identifier_not_position(self):
        payload = {"items": [self.item("B0SECOND01", "5000000000002"), self.item("B0FIRST001", "5000000000001"),
                             self.item("B0FIRST002", "5000000000001"),
                             self.item("B0UPC00001", "0019659200341")]}
        with patch.object(client_module.requests, "get", return_value=_response(payload=payload)) as get:
            res = self.client.search_catalog_items_by_identifier(
                ["5000000000001", "5000000000002", "5000000000003", "0019659200341"])
        params = get.call_args.kwargs["params"]
        self.assertEqual((params["identifiersType"], params["includedData"]), ("EAN", "identifiers,summaries"))
        self.assertEqual([m["asin"] for m in res["5000000000001"]], ["B0FIRST001", "B0FIRST002"])
        self.assertEqual(res["5000000000002"], [{"asin": "B0SECOND01", "title": "Title B0SECOND01", "brand": "Philips"}])
        self.assertEqual(res["5000000000003"], [])
        self.assertEqual([m["asin"] for m in res["0019659200341"]], ["B0UPC00001"])

    def test_failures_are_none_not_empty(self):
        with patch.object(client_module.requests, "get", side_effect=[_response(429)] * client_module.MAX_RETRIES):
            self.assertIsNone(self.client.search_catalog_items_by_identifier(["5000000000001"]))
        with patch.object(client_module.requests, "get", return_value=_response(400)):
            self.assertIsNone(self.client.search_catalog_items_by_identifier(["5000000000001"]))
        self.assertIsNone(self.client.search_catalog_items_by_identifier(["5000000000001"], marketplace="MOON"))

    def test_rate_limit_is_retried(self):
        ok = _response(payload={"items": [self.item("B0A0000001", "5000000000001")]})
        with patch.object(client_module.requests, "get", side_effect=[_response(429), ok]) as get:
            res = self.client.search_catalog_items_by_identifier(["5000000000001"])
        self.assertEqual(get.call_count, 2)
        self.assertEqual(res["5000000000001"][0]["asin"], "B0A0000001")


if __name__ == "__main__":
    unittest.main()
